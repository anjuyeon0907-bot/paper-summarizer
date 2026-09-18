import io
import os
import re
from xml.sax.saxutils import escape

import arxiv
import streamlit as st
from docx import Document
from dotenv import load_dotenv
from openai import OpenAI
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

load_dotenv()

st.set_page_config(page_title="논문 요약 도우미", page_icon=":material/science:")

NAMING_NOTE = (
    "학명, 화합물명, 시약명, 세포주명, 유전자명 등 고유명사는 임의로 번역하거나 다른 명칭으로 "
    "바꾸지 말고 논문에 쓰인 원문 표기를 그대로 사용하세요. 한국어 통용명이 논문에 없다면 "
    "지어내지 말고 원문(학명 등)을 그대로 쓰세요."
)

FORMULA_NOTE = (
    "수식은 LaTeX 문법(예: $...$, \\frac{}{}, \\text{})을 절대 쓰지 말고, 일반 텍스트로 풀어서 "
    "표기하세요. 예: 'LDH 방출률(%) = (실험군 흡광도 - 저대조군 흡광도) / (고대조군 흡광도 - "
    "저대조군 흡광도) × 100' 처럼 분수는 '/'로, 곱하기는 '×'로, 아래첨자·위첨자는 괄호나 텍스트로 풀어 쓰세요."
)

SUMMARY_PROMPT = f"""당신은 연구를 돕는 전문 분석가입니다. 아래 논문 내용을 분석하여 다음 5가지 항목을 Markdown 형식으로 한국어로 작성해주세요.
각 항목의 제목은 아래 형식을 그대로 사용하세요. {NAMING_NOTE} {FORMULA_NOTE}

## 🎯 핵심 요약 (3줄)

## 💡 주요 기여점 (Main Contribution)

## 🧪 제안하는 방법론 및 핵심 알고리즘

## ⚠️ 연구의 한계점 (Limitations)

## 🔍 내 연구에의 적용 포인트

---
논문 제목: {{title}}

논문 내용:
{{paper_text}}
"""

PROTOCOL_PROMPT = f"""당신은 실험 프로토콜을 정리하는 전문 연구보조원입니다. 아래 논문에서 Materials and Methods(재료 및 방법) 부분을 찾아,
다른 연구자가 실험을 그대로 재현할 수 있도록 Markdown으로 정리해주세요. {NAMING_NOTE} {FORMULA_NOTE}
논문에 명시되지 않은 내용은 추측해서 채우지 말고 "명시되지 않음"이라고 쓰세요.

## 🧫 재료 (시료·시약·세포주·장비)

## 🔬 실험 절차 (단계별)

## ⚙️ 주요 파라미터 (농도·온도·시간·배율 등)

## 📊 통계 분석 방법

---
논문 제목: {{title}}

논문 내용:
{{paper_text}}
"""

# Keep well under typical context limits regardless of model choice.
MAX_INPUT_CHARS = 15000

# Gemini exposes an OpenAI-compatible endpoint, so the openai client works as-is.
GEMINI_BASE_URL = "https://generativelanguage.googleapis.com/v1beta/openai/"

# Built-in CJK CID font: renders Korean in the PDF without bundling a font file.
pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
_PDF_BODY_STYLE = ParagraphStyle(name="Body", fontName="HYGothic-Medium", fontSize=10, leading=15)
_PDF_HEADING_STYLES = {
    level: ParagraphStyle(
        name=f"H{level}",
        fontName="HYGothic-Medium",
        fontSize=16 - level * 2,
        leading=20 - level * 2,
        spaceBefore=10,
        spaceAfter=6,
    )
    for level in (1, 2, 3, 4)
}

DOWNLOAD_FORMATS = ["Markdown (.md)", "TXT (.txt)", "Word (.docx)", "PDF (.pdf)"]


def slugify_filename(name: str, max_length: int = 60) -> str:
    name = re.sub(r'[\\/:*?"<>|]', "", name)
    name = re.sub(r"\s+", "_", name.strip())
    return name[:max_length] or "paper"


def markdown_to_plain_text(markdown_text: str) -> str:
    lines = []
    for line in markdown_text.splitlines():
        line = re.sub(r"^#{1,6}\s*", "", line)
        line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        line = re.sub(r"^-\s+", "• ", line)
        lines.append(line)
    return "\n".join(lines)


def markdown_to_docx_bytes(markdown_text: str) -> bytes:
    doc = Document()
    for line in markdown_text.splitlines():
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
            continue
        heading_match = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if heading_match:
            level = min(len(heading_match.group(1)), 4)
            doc.add_heading(heading_match.group(2), level=level)
        elif stripped.startswith("- "):
            doc.add_paragraph(re.sub(r"\*\*(.+?)\*\*", r"\1", stripped[2:]), style="List Bullet")
        else:
            doc.add_paragraph(re.sub(r"\*\*(.+?)\*\*", r"\1", stripped))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _md_inline_to_reportlab(text: str) -> str:
    text = escape(text)
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)


def markdown_to_pdf_bytes(markdown_text: str) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
    )
    story = []
    for line in markdown_text.splitlines():
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 6))
            continue
        heading_match = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if heading_match:
            level = min(len(heading_match.group(1)), 4)
            story.append(Paragraph(_md_inline_to_reportlab(heading_match.group(2)), _PDF_HEADING_STYLES[level]))
        elif stripped.startswith("- "):
            story.append(Paragraph(f"• {_md_inline_to_reportlab(stripped[2:])}", _PDF_BODY_STYLE))
        else:
            story.append(Paragraph(_md_inline_to_reportlab(stripped), _PDF_BODY_STYLE))
    doc.build(story)
    return buffer.getvalue()


def render_download_section(markdown_text: str, base_filename: str, key_prefix: str) -> None:
    fmt = st.selectbox(
        "다운로드 형식",
        DOWNLOAD_FORMATS,
        key=f"{key_prefix}_format",
    )
    if fmt == "Markdown (.md)":
        data, mime, ext = markdown_text.encode("utf-8"), "text/markdown", "md"
    elif fmt == "TXT (.txt)":
        data, mime, ext = markdown_to_plain_text(markdown_text).encode("utf-8"), "text/plain", "txt"
    elif fmt == "Word (.docx)":
        data = markdown_to_docx_bytes(markdown_text)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ext = "docx"
    else:
        data, mime, ext = markdown_to_pdf_bytes(markdown_text), "application/pdf", "pdf"

    st.download_button(
        "다운로드",
        data=data,
        file_name=f"{base_filename}.{ext}",
        mime=mime,
        icon=":material/download:",
        key=f"{key_prefix}_download",
    )


@st.cache_data(ttl="1h", max_entries=50)
def search_arxiv(query: str, max_results: int = 5) -> list[dict]:
    client = arxiv.Client()
    search = arxiv.Search(
        query=query,
        max_results=max_results,
        sort_by=arxiv.SortCriterion.SubmittedDate,
    )
    results = []
    for r in client.results(search):
        results.append(
            {
                "title": r.title,
                "authors": ", ".join(a.name for a in r.authors),
                "summary": r.summary,
                "url": r.entry_id,
                "published": r.published.strftime("%Y-%m-%d") if r.published else "",
            }
        )
    return results


def extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def run_llm(client: OpenAI, model: str, prompt_template: str, title: str, paper_text: str) -> str:
    prompt = prompt_template.format(title=title, paper_text=paper_text[:MAX_INPUT_CHARS])
    response = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.3,
    )
    return response.choices[0].message.content


with st.sidebar:
    st.header("설정")
    api_key = st.text_input(
        "Gemini API 키",
        type="password",
        value=os.getenv("GEMINI_API_KEY", ""),
        help="입력한 키는 저장되지 않고 이 세션에서만 사용됩니다. Google AI Studio에서 발급받은 키를 입력하세요.",
    )
    model = st.text_input(
        "모델",
        value="gemini-3.6-flash",
        help="Gemini 모델명은 자주 바뀝니다. 오류가 나면 https://ai.google.dev/gemini-api/docs/models 에서 현재 사용 가능한 모델명을 확인해 입력하세요.",
    )

st.title("논문 요약 도우미")
st.caption("arXiv 검색 또는 PDF 업로드로 논문을 분석하고 연구에 필요한 항목을 추출합니다.")

tab_search, tab_pdf, tab_protocol = st.tabs(
    [
        ":material/search: arXiv 검색",
        ":material/upload_file: PDF 업로드",
        ":material/biotech: 프로토콜 추출",
    ]
)

with tab_search:
    query = st.text_input("검색 키워드", placeholder="예: diffusion model")
    if st.button("검색", icon=":material/search:") and query:
        with st.spinner("arXiv에서 논문을 검색하는 중..."):
            st.session_state["arxiv_results"] = search_arxiv(query)

    for i, paper in enumerate(st.session_state.get("arxiv_results", [])):
        with st.container(border=True):
            st.markdown(f"**{paper['title']}**")
            st.caption(f"{paper['authors']} · {paper['published']}")
            with st.expander("초록 보기"):
                st.write(paper["summary"])

            if st.button("요약하기", key=f"summarize_arxiv_{i}", icon=":material/auto_awesome:"):
                if not api_key:
                    st.warning("사이드바에 API 키를 입력해주세요.")
                else:
                    with st.spinner("LLM으로 분석하는 중..."):
                        try:
                            client = OpenAI(api_key=api_key, base_url=GEMINI_BASE_URL)
                            st.session_state[f"summary_arxiv_{i}"] = run_llm(
                                client, model, SUMMARY_PROMPT, paper["title"], paper["summary"]
                            )
                        except Exception as e:
                            st.error(f"요약 중 오류가 발생했습니다: {e}")

            summary = st.session_state.get(f"summary_arxiv_{i}")
            if summary:
                st.markdown(summary)
                render_download_section(
                    summary,
                    f"{slugify_filename(paper['title'])}_summary",
                    key_prefix=f"arxiv_{i}",
                )

with tab_pdf:
    uploaded_file = st.file_uploader(
        "논문 PDF 업로드", type=["pdf"], key="pdf_uploader_summary"
    )

    if uploaded_file and st.button(
        "요약하기", key="summarize_pdf", icon=":material/auto_awesome:"
    ):
        if not api_key:
            st.warning("사이드바에 API 키를 입력해주세요.")
        else:
            with st.spinner("PDF에서 텍스트를 추출하는 중..."):
                pdf_text = extract_pdf_text(uploaded_file.read())
            with st.spinner("LLM으로 분석하는 중..."):
                try:
                    client = OpenAI(api_key=api_key, base_url=GEMINI_BASE_URL)
                    st.session_state["summary_pdf"] = run_llm(
                        client, model, SUMMARY_PROMPT, uploaded_file.name, pdf_text
                    )
                except Exception as e:
                    st.error(f"요약 중 오류가 발생했습니다: {e}")

    pdf_summary = st.session_state.get("summary_pdf")
    if pdf_summary and uploaded_file:
        st.markdown(pdf_summary)
        paper_name = os.path.splitext(uploaded_file.name)[0]
        render_download_section(
            pdf_summary,
            f"{slugify_filename(paper_name)}_summary",
            key_prefix="pdf_summary",
        )

with tab_protocol:
    st.caption("논문의 Materials and Methods를 바탕으로 재현 가능한 실험 프로토콜을 추출합니다.")
    protocol_file = st.file_uploader(
        "논문 PDF 업로드", type=["pdf"], key="pdf_uploader_protocol"
    )

    if protocol_file and st.button(
        "프로토콜 추출하기", key="extract_protocol", icon=":material/biotech:"
    ):
        if not api_key:
            st.warning("사이드바에 API 키를 입력해주세요.")
        else:
            with st.spinner("PDF에서 텍스트를 추출하는 중..."):
                protocol_text = extract_pdf_text(protocol_file.read())
            with st.spinner("LLM으로 프로토콜을 정리하는 중..."):
                try:
                    client = OpenAI(api_key=api_key, base_url=GEMINI_BASE_URL)
                    st.session_state["protocol_result"] = run_llm(
                        client, model, PROTOCOL_PROMPT, protocol_file.name, protocol_text
                    )
                except Exception as e:
                    st.error(f"프로토콜 추출 중 오류가 발생했습니다: {e}")

    protocol_result = st.session_state.get("protocol_result")
    if protocol_result and protocol_file:
        st.markdown(protocol_result)
        paper_name = os.path.splitext(protocol_file.name)[0]
        render_download_section(
            protocol_result,
            f"{slugify_filename(paper_name)}_protocol",
            key_prefix="protocol",
        )
